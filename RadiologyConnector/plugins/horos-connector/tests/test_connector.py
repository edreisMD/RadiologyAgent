import asyncio
import copy
import hashlib
import io
import json
import os
import subprocess
import sys
import shutil
from pathlib import Path
import pytest
from PIL import Image
from docx import Document
from horos_connector.engine import Horos, fingerprint
from horos_connector.queue import Queue
from horos_connector.media import export_study
from horos_connector.storage import write_json, component, contained
from horos_connector.reports import prepare_report, publish_report


def detail(uid="1.2.3", count=3, modality="CT"):
    return {"study":{"id":"study", "studyUID":uid, "patientName":"Example Patient", "patientID":"TEST-01", "title":"CT research fixture", "modality":modality, "date":0, "accession":"TEST"}, "frameCount":count, "series":[{"id":"series", "uid":uid+".1", "name":"Axial", "modality":modality, "frames":[{"id":f"frame-{i}", "sopInstanceUID":uid+"."+str(i+10), "index":i, "instance":i+1, "frame":0, "width":64, "height":64} for i in range(count)]}]}

class FakeEngine:
    def render(self, study_id, image_id, *args, **kwargs):
        i = int(image_id.split("-")[-1]); image = Image.new("RGB", (64,64), (i*30%255,100,140))
        return image, {"windowWidth":400,"windowCenter":40,"width":64,"height":64}

@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("RADAGENT_CONNECTOR_HOME", str(tmp_path/"private"))
    root=tmp_path/"private"; root.mkdir()
    return root


def test_new_arrival_waits_for_stable_inventory_without_historical_backfill(isolated):
    q=Queue(); q.observe([detail()], now=0); q.observe([detail()], now=100)
    assert not q.jobs()
    fresh=detail("1.2.4",2); q.observe([detail(),fresh],now=110)
    q.observe([detail(),fresh],now=125); assert not q.jobs()
    fresh=detail("1.2.4",3); q.observe([detail(),fresh],now=130)
    q.observe([detail(),fresh],now=155); assert not q.jobs()
    q.observe([detail(),fresh],now=161); assert len(q.jobs())==1
    q.observe([detail(),fresh],now=200); assert len(q.jobs())==1


def test_new_images_create_new_revision_and_claim_is_exclusive(isolated):
    q=Queue(); a=q.enqueue(detail()); b=q.enqueue(detail(count=4)); assert a!=b
    q.update(a,state="ready"); first=q.claim(a)
    with pytest.raises(ValueError): q.claim(a)
    with pytest.raises(ValueError): q.record_review(a,"wrong",[0])
    q.record_review(a,first["claim_token"],[0]); q.renew(a,first["claim_token"])
    assert q.get(a)["reviewed"]==[0]


def test_export_preserves_every_frame_and_video_count(isolated):
    result=export_study(FakeEngine(),detail())
    assert result["complete"] and result["rendered_frames"]==3
    series=result["series"][0]
    assert [f["index"] for f in series["frames"]]==[0,1,2]
    assert series["contact_sheets"][0]["image_indices"]==[0,1,2]
    video=series["videos"][0]["path"]
    probe=subprocess.check_output([shutil.which("ffprobe") or "/opt/homebrew/bin/ffprobe","-v","error","-select_streams","v:0","-count_frames","-show_entries","stream=nb_read_frames","-of","json",video])
    assert json.loads(probe)["streams"][0]["nb_read_frames"]=="3"
    class Fail:
        def render(self,*args): raise AssertionError("Cache should avoid rerendering")
    assert export_study(Fail(),detail())["complete"]


def test_xray_exports_original_rendered_pixels_without_cine(isolated):
    result=export_study(FakeEngine(),detail(count=1,modality="DX"))
    assert not result["series"][0]["videos"]
    assert Image.open(result["series"][0]["frames"][0]["path"]).getpixel((0,0))==(0,100,140)


def test_report_requires_review_and_preserves_user_edited_destination(isolated,tmp_path):
    q=Queue(); job=q.enqueue(detail()); q.update(job,state="ready"); token=q.claim(job)["claim_token"]
    with pytest.raises(ValueError): prepare_report(q,job,token,"FINDINGS\nExample research finding.")
    q.record_review(job,token,[0,1,2]); dest=tmp_path/"drive"; dest.mkdir()
    write_json(isolated/"config.json",{"report_directory":str(dest)})
    report=prepare_report(q,job,token,"FINDINGS\nExample research finding.\n\nIMPRESSION\nEvaluation only.")
    repeat=prepare_report(q,job,token,"FINDINGS\nExample research finding.\n\nIMPRESSION\nEvaluation only.")
    assert repeat["sha256"]==report["sha256"]
    doc=Document(report["path"]); assert doc.paragraphs[0].text=="Draft for evaluation"
    with pytest.raises(ValueError): publish_report(q,job,token,report["draft_id"],False)
    result=publish_report(q,job,token,report["draft_id"],True)
    path=Path(result["path"]); assert path.parent.name=="Example Patient - TEST-01"
    assert hashlib.sha256(path.read_bytes()).hexdigest()==report["sha256"]
    assert q.get(job)["state"]=="drafted"
    with pytest.raises(ValueError): q.claim(job)
    # A crash after writing the file but before completing the queue must never overwrite an edit.
    q.update(job,state="ready"); token=q.claim(job)["claim_token"]
    path.write_bytes(b"Clinician edited this report")
    with pytest.raises(ValueError): publish_report(q,job,token,report["draft_id"],True)
    assert path.read_bytes()==b"Clinician edited this report"


def test_partial_review_is_explicit_in_word(isolated):
    q=Queue(); job=q.enqueue(detail()); q.update(job,state="ready"); token=q.claim(job)["claim_token"]
    report=prepare_report(q,job,token,"FINDINGS\nNo interpretation made.","Images require radiologist review.")
    assert "Partial image review: 0 of 3" in "\n".join(p.text for p in Document(report["path"]).paragraphs)


def test_identity_paths_and_loopback_descriptor_validation(isolated,tmp_path):
    assert "/" not in component("../../Patient/Name")
    with pytest.raises(ValueError): contained(tmp_path, tmp_path/".."/"escape")
    path=tmp_path/"connection.json"
    write_json(path,{"port":443,"token":"x"*40,"protocolVersion":1,"pid":os.getpid()})
    with pytest.raises(ValueError): Horos(path).call("/health")
    with pytest.raises(ValueError): Horos(path).call("/arbitrary")


def test_export_claim_recovers_only_after_timeout(isolated):
    q=Queue(); job=q.enqueue(detail()); token=q.begin_export(job)
    assert token and q.begin_export(job) is None
    q.finish_export(job,"wrong"); assert q.get(job)["state"]=="exporting"
    q.finish_export(job,token); assert q.get(job)["state"]=="ready"


def test_protocol_image_result_and_bounds(isolated):
    from horos_connector.server import view_frames, view_contact_sheet
    q=Queue(); d=detail(count=2,modality="DX"); job=q.enqueue(d)
    export_study(FakeEngine(),d); q.update(job,state="ready"); token=q.claim(job)["claim_token"]
    content=view_frames(job,[0,1],token).content
    assert [c.type for c in content]==["text","image","text","image"]
    assert q.get(job)["reviewed"]==[0,1]
    with pytest.raises(ValueError): view_frames(job,[999],token)
    assert view_contact_sheet(job,0).content[1].type=="image"


def test_stdio_sdk_handshake_and_tool_discovery(isolated):
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    async def run():
        params=StdioServerParameters(command=sys.executable,args=["-m","horos_connector.server"],env=dict(os.environ))
        async with stdio_client(params) as (read,write):
            async with ClientSession(read,write) as client:
                await client.initialize(); tools=await client.list_tools()
                names={t.name for t in tools.tools}
                assert {"find_studies","view_frames","prepare_report","publish_report","claim_study","open_workspace","inspect_view","current_view","update_workspace_report","reserve_study_run","attach_study_run"}<=names
                assert not any("finalize" in n for n in names)
                result=await client.call_tool("worklist",{})
                assert not result.isError
    asyncio.run(run())

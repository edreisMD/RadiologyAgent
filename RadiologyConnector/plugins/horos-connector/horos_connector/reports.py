from __future__ import annotations
import hashlib
import json
import os
import re
import tempfile
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .storage import component, config, contained, data_root, digest, write_json


def build_report(study: dict, document: str, output: Path, coverage_note=""):
    from .privacy import public_study
    study=public_study(study)
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"]:
        style = doc.styles[name]; style.font.name = "Arial"; style.font.color.rgb = RGBColor(0,0,0)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(5)
        if style.element.pPr is not None:
            for node in list(style.element.pPr):
                if node.tag in {qn("w:pBdr"), qn("w:shd")}: style.element.pPr.remove(node)
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Title"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].font.size = Pt(10.5); doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.add_paragraph("Draft for evaluation", "Title")
    doc.add_paragraph("Not for medical use, research only.", "Subtitle")
    p = doc.add_paragraph(); p.add_run(study.get("patientName") or "Patient").bold = True
    date = datetime.fromtimestamp(study["date"]).strftime("%d/%m/%Y") if study.get("date") else "Not provided"
    doc.add_paragraph(f"Patient ID: {study.get('patientID') or 'Not provided'}    Study date: {date}")
    doc.add_paragraph(f"{study.get('title') or study.get('modality') or 'Imaging study'}    Accession: {study.get('accession') or 'Not provided'}")
    for line in document.splitlines():
        clean = re.sub(r'^\s*#{1,6}\s+', '', line.strip()).replace("**", "")
        if not clean: continue
        if clean.upper().startswith(("DRAFT FOR EVALUATION", "RASCUNHO PARA AVALIAÇÃO")): continue
        heading = len(clean) < 85 and any(c.isalpha() for c in clean) and clean == clean.upper()
        doc.add_paragraph(clean, "Heading 1" if heading else "Normal")
    if coverage_note:
        doc.add_paragraph("Review coverage", "Heading 1"); doc.add_paragraph(coverage_note)
    footer = section.footer.paragraphs[0]
    footer.text = "DRAFT FOR EVALUATION • Not for medical use, research only. • "
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    footer.style = doc.styles["Footer"]; footer.paragraph_format.space_after = Pt(0)
    doc.core_properties.author = "Radiology Agent"
    doc.core_properties.title = "Draft for evaluation"
    doc.core_properties.subject = study.get("title", "Radiology report")
    doc.core_properties.comments = "Unsigned research evaluation draft"
    output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    doc.save(output); os.chmod(output, 0o600)


def prepare_report(queue, job_id, token, document: str, limitations=""):
    job = queue.authorize(job_id, token)
    if not document.strip() or len(document) > 60_000: raise ValueError("Report document is empty or exceeds 60,000 characters.")
    all_indices = {f["index"] for s in job["detail"]["series"] for f in s["frames"]}
    missing = all_indices - set(job["reviewed"])
    if missing and not limitations.strip():
        raise ValueError(f"{len(missing)} frames have not been returned as full-detail images. Review them or explicitly describe the limited review in limitations.")
    review_text = limitations.strip()
    if missing: review_text = f"Partial image review: {len(all_indices)-len(missing)} of {len(all_indices)} available frames inspected at full detail. " + review_text
    version = digest([job_id, document, review_text])[:20]
    output = data_root() / "reports" / job_id / f"{version}.docx"
    if output.exists() and output.with_suffix(".json").exists():
        previous = json.loads(output.with_suffix(".json").read_text())
        if previous["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest():
            return {**previous, "next_step": "Render and visually inspect this DOCX, then publish_report with layout_reviewed=true."}
    build_report(job["detail"]["study"], document, output, review_text)
    content_hash = hashlib.sha256(output.read_bytes()).hexdigest()
    receipt = {"job_id": job_id, "draft_id": version, "path": str(output), "sha256": content_hash, "reviewed_frames": len(job["reviewed"]), "total_frames": len(all_indices), "purpose": "Draft for evaluation"}
    write_json(output.with_suffix(".json"), receipt)
    return {**receipt, "next_step": "Render and visually inspect this DOCX, then publish_report with layout_reviewed=true."}


def publish_report(queue, job_id, token, draft_id, layout_reviewed):
    if not layout_reviewed: raise ValueError("Render and visually inspect the report before publication.")
    if not re.fullmatch(r"[0-9a-f]{20}", draft_id): raise ValueError("Invalid draft ID")
    job = queue.authorize(job_id, token)
    source = data_root() / "reports" / job_id / f"{draft_id}.docx"
    receipt = json.loads(source.with_suffix(".json").read_text())
    raw = source.read_bytes()
    if receipt["sha256"] != hashlib.sha256(raw).hexdigest(): raise ValueError("Draft changed after preparation. Prepare the corrected document again.")
    configured = config().get("report_directory")
    if not configured: raise ValueError("Set the report directory in the connector configuration first.")
    root = Path(configured).expanduser()
    if not root.is_dir(): raise ValueError("The configured report folder is unavailable. Reconnect Google Drive; the draft remains local.")
    from .privacy import public_study
    study = public_study(job["detail"]["study"])
    folder = contained(root, root / (component(study.get("patientName", "")) + " - " + component(study.get("patientID", ""), digest(study["studyUID"])[:10])))
    folder.mkdir(exist_ok=True, mode=0o700)
    date = datetime.fromtimestamp(study["date"]).strftime("%Y-%m-%d") if study.get("date") else "Undated"
    name = f"{date} - {component(study.get('title', ''), 'Study', 65)} - Draft for evaluation - {job_id[:8]}-{draft_id[:8]}.docx"
    destination = contained(root, folder / name)
    # Link an atomic temp file into place without replacing any clinician-edited report.
    if destination.exists():
        if hashlib.sha256(destination.read_bytes()).hexdigest() != receipt["sha256"]:
            raise ValueError("A report at this destination has been edited. It was preserved.")
    else:
        fd, temporary = tempfile.mkstemp(dir=folder, prefix=".radagent-")
        try:
            with os.fdopen(fd, "wb") as out: out.write(raw); out.flush(); os.fsync(out.fileno())
            try: os.link(temporary, destination)
            except FileExistsError:
                if hashlib.sha256(destination.read_bytes()).hexdigest() != receipt["sha256"]: raise ValueError("Destination changed during publication.")
        finally: os.unlink(temporary)
    if hashlib.sha256(destination.read_bytes()).hexdigest() != receipt["sha256"]: raise RuntimeError("Report verification failed.")
    queue.complete(job_id, token, destination)
    return {"saved": True, "path": str(destination), "status": "Draft for evaluation", "cloud_sync": "Written to the configured folder; Google Drive controls cloud synchronization."}
